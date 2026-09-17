import base64
import io
import json
import logging
import os
import re
import tempfile
from pathlib import Path

from fastapi import FastAPI
from fastapi.responses import FileResponse
from pydantic import BaseModel
from openai import OpenAI
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

app = FastAPI()
logger = logging.getLogger("uvicorn.error")


class AnalysisRequest(BaseModel):
    query: str
    language: str = "fa"


class VisualRequest(BaseModel):
    query: str
    analysis: str = ""
    language: str = "fa"


class ImageRequest(BaseModel):
    query: str
    analysis: str = ""
    kind: str = "infographic"
    language: str = "fa"


class ExportRequest(BaseModel):
    query: str
    analysis: str = ""
    visuals: dict = {}
    format: str = "pdf"


class PublishRequest(BaseModel):
    query: str
    analysis: str = ""
    visuals: dict = {}


def get_openai_key():
    """Find a configured OpenAI key without exposing its value."""
    candidates = ["OPENAI_API_KEY", "OPENAI_API_KEY_2", "OPENAI_API_KEY_OLD"]
    for name in candidates:
        value = os.getenv(name)
        if value and value.strip():
            return value.strip(), name
    for name, value in os.environ.items():
        if name.startswith("OPENAI_API_KEY_") and value and value.strip():
            return value.strip(), name
    return None, None


def get_model():
    return os.getenv("OPENAI_MODEL", "gpt-5.6-luna").strip() or "gpt-5.6-luna"


@app.on_event("startup")
def startup_diagnostics():
    key, key_name = get_openai_key()
    model = get_model()
    openai_names = sorted(name for name in os.environ if name.startswith("OPENAI_API_KEY"))
    logger.info("OPENAI_RUNTIME_DIAGNOSTICS configured=%s source=%s key_length=%s variables=%s model=%s", bool(key), key_name or "NONE", len(key) if key else 0, openai_names, model)


@app.get("/")
def home():
    return FileResponse("dashboard.html")


@app.get("/api")
def api_status():
    key, key_name = get_openai_key()
    return {"status": "online", "system": "AI Strategic Studio", "openai_configured": bool(key), "openai_key_source": key_name, "model": get_model(), "image_model": "gpt-image-2"}


@app.get("/api/diagnostics")
def diagnostics():
    key, key_name = get_openai_key()
    openai_names = sorted(name for name in os.environ if name.startswith("OPENAI_API_KEY"))
    return {"ok": bool(key), "key_source": key_name, "key_length": len(key) if key else 0, "key_variables_present": openai_names, "model": get_model(), "image_model": "gpt-image-2"}


@app.post("/api/analyze")
def analyze(request: AnalysisRequest):
    q = request.query.strip()
    if not q:
        return {"status": "error", "message": "موضوع تحلیل وارد نشده است."}
    api_key, key_name = get_openai_key()
    if not api_key:
        return {"status": "error", "message": "هیچ متغیر محیطی معتبر برای کلید OpenAI در زمان اجرا پیدا نشد.", "hint": "در Render باید OPENAI_API_KEY روی همین Web Service در runtime موجود باشد."}
    try:
        model = get_model()
        client = OpenAI(api_key=api_key)
        lang = (request.language or "fa").lower()
        lang_name = {"fa": "فارسی", "ar": "العربية", "en": "English"}.get(lang, "فارسی")
        response = client.responses.create(
            model=model,
            instructions=(f"تو موتور تحلیل راهبردی AI Strategic Studio هستی. پاسخ را به زبان {lang_name} و به‌صورت بی‌طرف و تحلیلی ارائه کن. "
                          "واقعیت، تحلیل و عدم قطعیت را از هم جدا کن و از ساختن منبع یا داده خودداری کن. "
                          "موضوع را در چهار بخش خلاصه، سیاسی، اقتصادی و امنیتی بررسی کن."),
            input=q,
        )
        return {"status": "ok", "query": q, "analysis": {"summary": response.output_text, "political": "تحلیل سیاسی در متن بالا ارائه شده است.", "economic": "تحلیل اقتصادی در متن بالا ارائه شده است.", "security": "تحلیل امنیتی در متن بالا ارائه شده است.", "next_step": "موتور AI فعال است."}}
    except Exception as e:
        error_text = str(e).strip()
        return {"status": "error", "message": "خطا در موتور AI: " + (error_text[:500] if error_text else "خطای نامشخص"), "model": model, "key_source": key_name}


@app.post("/api/generate-visuals")
def generate_visuals(request: VisualRequest):
    q = request.query.strip()
    if not q:
        return {"status": "error", "message": "موضوع تحلیل وارد نشده است."}
    api_key, key_name = get_openai_key()
    if not api_key:
        return {"status": "error", "message": "کلید OpenAI در runtime تنظیم نشده است."}
    lang = (request.language or "fa").lower()
    lang_name = {"fa": "فارسی", "ar": "العربية", "en": "English"}.get(lang, "فارسی")
    prompt = f"""
برای داشبورد AI Strategic Studio، برای موضوع زیر داده بصری ساختاریافته تولید کن.
زبان خروجی همه عنوان‌ها، برچسب‌ها و یادداشت‌ها: {lang_name}
موضوع: {q}
متن تحلیل موجود (ممکن است خالی باشد): {request.analysis[:9000]}

فقط JSON معتبر برگردان و هیچ Markdown یا توضیح بیرونی نده.
ساختار دقیق:
{{
  "title": "عنوان کوتاه فارسی",
  "map": {{"center": "middle_east", "locations": [{{"city":"نام شهر شناخته‌شده", "category":"political|economic|security|general", "note":"توضیح کوتاه فارسی"}}]}},
  "infographic": {{"title":"عنوان کوتاه فارسی", "dimensions":[{{"label":"سیاسی|اقتصادی|امنیتی|اجتماعی|انرژی|دیپلماسی", "score":0, "note":"توضیح کوتاه"}}]}}
}}
قواعد: نقشه تحلیلی و غیرعملیاتی باشد؛ فقط شهرهای بزرگ و شناخته‌شده را نام ببر؛ مختصات تولید نکن؛ حداکثر 8 شهر و 6 بُعد؛ score ارزیابی کیفی مدل بین 0 تا 100 است و آمار واقعی نیست؛ از ادعا درباره کنترل سرزمینی، موقعیت نیروها، هدف‌گیری یا دستورالعمل عملیاتی خودداری کن.
"""
    try:
        model = get_model()
        client = OpenAI(api_key=api_key)
        response = client.responses.create(model=model, instructions="فقط JSON معتبر و سازگار با ساختار درخواست برگردان.", input=prompt)
        raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", response.output_text.strip(), flags=re.IGNORECASE | re.DOTALL).strip()
        data = json.loads(raw)
        return {"status": "ok", "query": q, "visuals": data}
    except Exception as e:
        logger.exception("VISUAL_GENERATION_ERROR")
        return {"status": "error", "message": "تولید داده بصری ناموفق بود: " + (str(e).strip()[:500] or "خطای نامشخص"), "model": get_model(), "key_source": key_name}


@app.post("/api/generate-image")
def generate_image(request: ImageRequest):
    q = request.query.strip()
    if not q:
        return {"status": "error", "message": "موضوع تحلیل وارد نشده است."}
    api_key, key_name = get_openai_key()
    if not api_key:
        return {"status": "error", "message": "کلید OpenAI در runtime تنظیم نشده است."}
    kind = request.kind.strip().lower() or "infographic"
    if kind not in {"infographic", "map"}:
        kind = "infographic"
    if kind == "map":
        visual_prompt = "یک نقشه تصویری بسیار تمیز و مینیمال از خاورمیانه برای یک گزارش تحلیلی بساز؛ شهرهای اصلی را معقول نمایش بده، بدون عملیات نظامی، اهداف، مسیر حمله، استقرار نیرو یا اطلاعات تاکتیکی. تصویر صرفاً illustrative است."
    else:
        visual_prompt = "یک اینفوگرافیک حرفه‌ای و مدرن برای یک اتاق فکر راهبردی بساز؛ نمودارهای ساده، کارت‌های شاخص و تایپوگرافی خوانا؛ از اعداد ساختگی به‌عنوان آمار واقعی استفاده نکن؛ تمرکز روی روندها، روابط و پیامدهای تحلیلی باشد."
    lang = (request.language or "fa").lower()
    lang_name = {"fa": "Persian", "ar": "Arabic", "en": "English"}.get(lang, "Persian")
    prompt = f"{visual_prompt}\nزبان متن‌های داخل تصویر: {lang_name}\nموضوع گزارش: {q}\nخلاصه تحلیل: {request.analysis[:5000]}\nسبک: dark professional intelligence dashboard, clean layout, editorial quality."
    try:
        client = OpenAI(api_key=api_key)
        result = client.images.generate(model="gpt-image-2", prompt=prompt, size="1536x1024")
        item = result.data[0]
        b64 = getattr(item, "b64_json", None)
        if not b64:
            return {"status": "error", "message": "مدل تصویر پاسخ قابل نمایش برنگرداند.", "image_model": "gpt-image-2"}
        return {"status": "ok", "kind": kind, "image_model": "gpt-image-2", "mime": "image/png", "image_data": "data:image/png;base64," + base64.b64encode(base64.b64decode(b64)).decode("ascii")}
    except Exception as e:
        logger.exception("IMAGE_GENERATION_ERROR")
        return {"status": "error", "message": "تولید تصویر ناموفق بود: " + (str(e).strip()[:500] or "خطای نامشخص"), "image_model": "gpt-image-2", "key_source": key_name}


CITY_COORDS = {
    "تهران": (35.689, 51.389), "ریاض": (24.714, 46.675), "آنکارا": (39.933, 32.86), "استانبول": (41.008, 28.978),
    "قاهره": (30.044, 31.236), "بغداد": (33.315, 44.366), "دمشق": (33.514, 36.277), "بیروت": (33.894, 35.502),
    "صنعا": (15.369, 44.191), "مسقط": (23.588, 58.383), "دوحه": (25.285, 51.531), "ابوظبی": (24.454, 54.377),
    "کویت": (29.376, 47.977), "عمان": (31.954, 35.911), "اورشلیم": (31.768, 35.214), "تل‌آویو": (32.085, 34.782),
    "Tehran": (35.689, 51.389), "Riyadh": (24.714, 46.675), "Cairo": (30.044, 31.236), "Baghdad": (33.315, 44.366),
    "Damascus": (33.514, 36.277), "Beirut": (33.894, 35.502), "Sanaa": (15.369, 44.191)
}


def make_map_infographic(visuals, path):
    W, H = 1600, 900
    img = Image.new("RGB", (W, H), (7, 17, 31))
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 34)
        small = ImageFont.truetype("DejaVuSans.ttf", 22)
    except Exception:
        font = small = ImageFont.load_default()
    draw.text((W-40, 35), "AI Strategic Studio — نقشه اینفوگرافیک تحلیلی", anchor="ra", font=font, fill=(220,235,255))
    # Approximate Middle East frame, used only as a visual briefing diagram.
    x0, y0, x1, y1 = 130, 150, 1470, 790
    draw.rounded_rectangle((x0, y0, x1, y1), radius=28, outline=(40, 65, 95), width=3, fill=(10, 25, 42))
    for i in range(1, 5):
        yy = y0 + i*(y1-y0)//5
        draw.line((x0, yy, x1, yy), fill=(22, 43, 66), width=1)
    for i in range(1, 6):
        xx = x0 + i*(x1-x0)//6
        draw.line((xx, y0, xx, y1), fill=(22, 43, 66), width=1)
    locations = (visuals or {}).get("map", {}).get("locations", []) if isinstance(visuals, dict) else []
    pts = []
    for item in locations[:8]:
        city = str(item.get("city", ""))
        if city not in CITY_COORDS:
            continue
        lat, lon = CITY_COORDS[city]
        px = int(x0 + (lon-25)/(60-25)*(x1-x0))
        py = int(y1 - (lat-12)/(45-12)*(y1-y0))
        cat = item.get("category", "general")
        fill = {"political": (96,165,250), "economic": (52,211,153), "security": (251,191,36), "general": (196,181,253)}.get(cat, (196,181,253))
        draw.ellipse((px-13, py-13, px+13, py+13), fill=fill, outline=(235,245,255), width=2)
        draw.text((px+18, py-10), city, font=small, fill=(235,245,255))
        pts.append((px, py))
    draw.text((W-50, H-55), "نمای اینفوگرافیک؛ جایگزین نقشه مرجع یا داده عملیاتی نیست", anchor="rs", font=small, fill=(142,163,188))
    img.save(path, format="PNG")


def clean_text(text):
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text or "")
    text = re.sub(r"^#{1,3}\s*", "", text, flags=re.MULTILINE)
    return text.strip()


def build_report_files(req):
    q = req.query.strip()
    visuals = req.visuals or {}
    analysis = clean_text(req.analysis)
    tmp = Path(tempfile.mkdtemp(prefix="strategic_report_"))
    map_png = tmp / "map_infographic.png"
    make_map_infographic(visuals, map_png)
    fmt = req.format.lower().strip()
    if fmt not in {"pdf", "docx"}:
        raise ValueError("فرمت خروجی فقط PDF یا Word است.")
    if fmt == "pdf":
        out = tmp / "ai-strategic-report.pdf"
        styles = getSampleStyleSheet()
        rtl = ParagraphStyle("rtl", parent=styles["BodyText"], fontName="Helvetica", fontSize=10, leading=16, alignment=TA_RIGHT)
        title = ParagraphStyle("title", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, leading=24, alignment=TA_RIGHT)
        doc = SimpleDocTemplate(str(out), pagesize=A4, rightMargin=15*mm, leftMargin=15*mm, topMargin=15*mm, bottomMargin=15*mm)
        story = [Paragraph("AI Strategic Studio — گزارش تحلیلی", title), Spacer(1, 5*mm), Paragraph("موضوع: " + q, rtl), Spacer(1, 4*mm)]
        story.append(RLImage(str(map_png), width=170*mm, height=95*mm))
        story.append(Spacer(1, 5*mm))
        story.append(Paragraph("خلاصه تحلیل", title))
        for block in re.split(r"\n{2,}", analysis):
            if block.strip(): story.append(Paragraph(block.replace("\n", "<br/>"), rtl)); story.append(Spacer(1, 2*mm))
        dims = (visuals.get("infographic", {}) if isinstance(visuals, dict) else {}).get("dimensions", [])
        if dims:
            story.append(Spacer(1, 3*mm)); story.append(Paragraph("شاخص‌های اینفوگرافیک", title))
            data = [["بُعد", "امتیاز کیفی", "یادداشت"]] + [[str(x.get("label","")), str(x.get("score","")), str(x.get("note",""))] for x in dims[:6]]
            table = Table(data, colWidths=[35*mm, 30*mm, 110*mm])
            table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#17335b")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), 0.5, colors.grey), ("ALIGN", (0,0), (-1,-1), "RIGHT"), ("FONTSIZE", (0,0), (-1,-1), 8), ("VALIGN", (0,0), (-1,-1), "TOP")]))
            story.append(table)
        doc.build(story)
    else:
        out = tmp / "ai-strategic-report.docx"
        doc = Document()
        doc.add_heading("AI Strategic Studio — گزارش تحلیلی", 0)
        p = doc.add_paragraph(); p.add_run("موضوع: ").bold = True; p.add_run(q)
        doc.add_picture(str(map_png), width=Inches(6.4))
        doc.add_heading("خلاصه تحلیل", level=1)
        for block in re.split(r"\n{2,}", analysis):
            if block.strip(): doc.add_paragraph(block.strip())
        dims = (visuals.get("infographic", {}) if isinstance(visuals, dict) else {}).get("dimensions", [])
        if dims:
            doc.add_heading("شاخص‌های اینفوگرافیک", level=1)
            table = doc.add_table(rows=1, cols=3)
            for i, h in enumerate(["بُعد", "امتیاز کیفی", "یادداشت"]): table.rows[0].cells[i].text = h
            for x in dims[:6]:
                cells = table.add_row().cells
                cells[0].text = str(x.get("label", "")); cells[1].text = str(x.get("score", "")); cells[2].text = str(x.get("note", ""))
        doc.save(out)
    data = base64.b64encode(out.read_bytes()).decode("ascii")
    mime = "application/pdf" if fmt == "pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return {"status":"ok", "filename":out.name, "mime":mime, "file_data":"data:"+mime+";base64,"+data}


@app.post("/api/export-report")
def export_report(request: ExportRequest):
    try:
        return build_report_files(request)
    except Exception as e:
        logger.exception("REPORT_EXPORT_ERROR")
        return {"status":"error", "message":"ساخت گزارش ناموفق بود: "+(str(e).strip()[:500] or "خطای نامشخص")}


@app.post("/api/publish-text")
def publish_text(request: PublishRequest):
    q = request.query.strip()
    analysis = clean_text(request.analysis)
    visuals = request.visuals or {}
    dims = (visuals.get("infographic", {}) if isinstance(visuals, dict) else {}).get("dimensions", [])
    lines = ["📌 گزارش تحلیلی — AI Strategic Studio", "", "موضوع: " + q, "", "🧠 جمع‌بندی", analysis[:6500]]
    if dims:
        lines += ["", "📊 شاخص‌های کیفی"]
        lines += [f"• {x.get('label','')}: {x.get('score','')} — {x.get('note','')}" for x in dims[:6]]
    lines += ["", "🗺️ نقشه گزارش: به‌صورت اینفوگرافیک در نسخه PDF و Word درج شده است.", "", "منبع/داده‌های واقعی باید پیش از انتشار راستی‌آزمایی شوند."]
    return {"status":"ok", "text":"\n".join(lines)}
