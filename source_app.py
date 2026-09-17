import base64
import html
import json
import logging
import os
import re
import tempfile
from pathlib import Path

from fastapi import FastAPI
from fastapi.responses import HTMLResponse
from pydantic import BaseModel
from openai import OpenAI
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches

from app import app as core_app, AnalysisRequest, ExportRequest, PublishRequest, analyze as core_analyze, make_map_infographic, clean_text
from source_catalog import SOURCE_GROUPS, ALL_DOMAINS, source_groups_for_query

logger = logging.getLogger("uvicorn.error")
app = FastAPI(title="AI Strategic Studio Research Gateway")


class SourceSearchRequest(BaseModel):
    query: str
    group: str = "همه"
    limit: int = 10


class SourceAnalyzeRequest(BaseModel):
    query: str
    analysis: str = ""


def get_openai_key():
    for name in ("OPENAI_API_KEY", "OPENAI_API_KEY_2", "OPENAI_API_KEY_OLD"):
        value = os.getenv(name)
        if value and value.strip():
            return value.strip()
    for name, value in os.environ.items():
        if name.startswith("OPENAI_API_KEY_") and value and value.strip():
            return value.strip()
    return None


def _dump(obj):
    try:
        return obj.model_dump()
    except Exception:
        try:
            return json.loads(obj.model_dump_json())
        except Exception:
            return None


def _collect_sources(node, out):
    if isinstance(node, dict):
        if isinstance(node.get("sources"), list):
            for src in node["sources"]:
                if isinstance(src, dict):
                    url = src.get("url") or src.get("link")
                    title = src.get("title") or src.get("name") or url
                    if url:
                        out.append({"title": str(title), "url": str(url)})
        for value in node.values():
            _collect_sources(value, out)
    elif isinstance(node, list):
        for value in node:
            _collect_sources(value, out)


def search_web(query, group="همه", limit=10):
    key = get_openai_key()
    if not key:
        return {"status": "error", "message": "کلید OpenAI در runtime تنظیم نشده است.", "sources": []}
    query = query.strip()
    if not query:
        return {"status": "error", "message": "عبارت جستجو وارد نشده است.", "sources": []}
    groups = source_groups_for_query(group)
    group_domains = sorted({d for g in groups.values() for d in g["domains"]})
    source_instruction = ""
    if group != "همه" and group in groups:
        source_instruction = f"فقط از منابع مرتبط با گروه «{groups[group]['label']}» استفاده کن."
    else:
        source_instruction = "در صورت امکان، ترکیبی از منابع رسمی، دانشگاهی، اندیشکده‌ای، خبری عربی و اسرائیلی و منابع غیررسمی معتبر ارائه کن."
    prompt = f"""
برای موضوع زیر یک جستجوی پژوهشی چندمنبعی انجام بده: {query}
{source_instruction}
منابع رسمی را از غیررسمی جدا کن. برای هر منبع عنوان و URL بده. ادعای خبری یا تحلیلی را به منبع نسبت بده و از ساختن URL خودداری کن.
اگر منبع غیررسمی است صریحاً آن را «غیررسمی» برچسب بزن.
حداکثر {max(3, min(limit, 15))} منبع مرتبط انتخاب کن.
"""
    try:
        client = OpenAI(api_key=key)
        tool = {"type": "web_search"}
        if group != "همه" and group_domains:
            tool["filters"] = {"allowed_domains": group_domains}
        response = client.responses.create(
            model=os.getenv("OPENAI_MODEL", "gpt-5.6-luna").strip() or "gpt-5.6-luna",
            tools=[tool],
            include=["web_search_call.action.sources"],
            input=prompt,
        )
        raw_sources = []
        _collect_sources(_dump(response) or {}, raw_sources)
        seen = set()
        sources = []
        for src in raw_sources:
            url = src["url"]
            if url in seen:
                continue
            seen.add(url)
            matched_group = "منبع وب"
            host = re.sub(r"^https?://(www\.)?", "", url).split("/")[0].lower()
            for name, data in SOURCE_GROUPS.items():
                if any(host.endswith(d) for d in data["domains"]):
                    matched_group = data["label"]
                    break
            sources.append({**src, "group": matched_group, "officiality": "غیررسمی" if group == "غیررسمی" or host in {"substack.com", "medium.com", "foreignpolicy.com", "responsiblestatecraft.org", "al-monitor.com"} else "رسمی/سازمانی یا رسانه‌ای"})
            if len(sources) >= limit:
                break
        return {"status": "ok", "query": query, "sources": sources, "summary": response.output_text}
    except Exception as exc:
        logger.exception("SOURCE_SEARCH_ERROR")
        return {"status": "error", "message": "جستجوی منابع ناموفق بود: " + (str(exc).strip()[:500] or "خطای نامشخص"), "sources": []}


def sources_for_report(query, analysis=""):
    search = search_web(query, "همه", 10)
    return search.get("sources", []) if search.get("status") == "ok" else []


def source_block(sources):
    if not sources:
        return ""
    lines = ["", "منابع مرتبط و قابل بررسی:"]
    for i, s in enumerate(sources[:10], 1):
        lines.append(f"{i}. {s.get('title','')} — {s.get('group','')} — {s.get('officiality','')} — {s.get('url','')}")
    return "\n".join(lines)


def rtl_ready(text):
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(text or ""))
    except Exception:
        return text or ""


def build_source_report(req, sources):
    q = req.query.strip()
    analysis = clean_text(req.analysis)
    lang = (getattr(req, "language", "fa") or "fa").lower()
    labels = {"fa": {"title":"AI Strategic Studio — گزارش تحلیلی","topic":"موضوع","analysis":"تحلیل","dims":"شاخص‌های تحلیلی","sources":"منابع مرتبط"}, "ar": {"title":"AI Strategic Studio — تقرير تحليلي","topic":"الموضوع","analysis":"التحليل","dims":"المؤشرات التحليلية","sources":"المصادر ذات الصلة"}, "en": {"title":"AI Strategic Studio — Analytical Report","topic":"Topic","analysis":"Analysis","dims":"Analytical Indicators","sources":"Related Sources"}}.get(lang, {"title":"AI Strategic Studio — گزارش تحلیلی","topic":"موضوع","analysis":"تحلیل","dims":"شاخص‌های تحلیلی","sources":"منابع مرتبط"})
    visuals = req.visuals or {}
    tmp = Path(tempfile.mkdtemp(prefix="strategic_report_sources_"))
    map_png = tmp / "map_infographic.png"
    make_map_infographic(visuals, map_png)
    fmt = req.format.lower().strip()
    if fmt not in {"pdf", "docx"}:
        raise ValueError("فرمت خروجی فقط PDF یا Word است.")
    if fmt == "pdf":
        out = tmp / "ai-strategic-report.pdf"
        font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
        if Path(font_path).exists():
            try: pdfmetrics.registerFont(TTFont("DejaVuSans", font_path))
            except Exception: pass
        font_name = "DejaVuSans" if "DejaVuSans" in pdfmetrics.getRegisteredFontNames() else "Helvetica"
        styles = getSampleStyleSheet()
        body = ParagraphStyle("body_rtl", parent=styles["BodyText"], fontName=font_name, fontSize=9.5, leading=15, alignment=TA_RIGHT)
        title = ParagraphStyle("title_rtl", parent=styles["Title"], fontName=font_name, fontSize=16, leading=22, alignment=TA_RIGHT)
        doc = SimpleDocTemplate(str(out), pagesize=A4, rightMargin=15*mm, leftMargin=15*mm, topMargin=15*mm, bottomMargin=15*mm)
        story = [Paragraph(rtl_ready(labels["title"]), title), Spacer(1, 4*mm), Paragraph(rtl_ready(labels["topic"] + ": " + q), body), Spacer(1, 4*mm), RLImage(str(map_png), width=170*mm, height=95*mm), Spacer(1, 4*mm), Paragraph(rtl_ready(labels["analysis"]), title)]
        for block in re.split(r"\n{2,}", analysis):
            if block.strip():
                story.append(Paragraph(rtl_ready(block.replace("\n", "<br/>")), body)); story.append(Spacer(1, 2*mm))
        dims = (visuals.get("infographic", {}) if isinstance(visuals, dict) else {}).get("dimensions", [])
        if dims:
            story += [Spacer(1, 3*mm), Paragraph(rtl_ready(labels["dims"]), title)]
            data = [[rtl_ready("بُعد"), rtl_ready("امتیاز کیفی"), rtl_ready("یادداشت")]] + [[rtl_ready(str(x.get("label",""))), str(x.get("score","")), rtl_ready(str(x.get("note","")))] for x in dims[:6]]
            table = Table(data, colWidths=[35*mm, 30*mm, 110*mm])
            table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#17335b")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), 0.5, colors.grey), ("ALIGN", (0,0), (-1,-1), "RIGHT"), ("FONTSIZE", (0,0), (-1,-1), 8), ("VALIGN", (0,0), (-1,-1), "TOP")]))
            story.append(table)
        if sources:
            story += [Spacer(1, 4*mm), Paragraph(rtl_ready(labels["sources"]), title)]
            for i, s in enumerate(sources[:10], 1):
                text = f"{i}. {s.get('title','')}<br/>{s.get('group','')} — {s.get('officiality','')}<br/>{s.get('url','')}"
                story += [Paragraph(text, body), Spacer(1, 2*mm)]
        doc.build(story)
    else:
        out = tmp / "ai-strategic-report.docx"
        doc = Document()
        doc.add_heading(labels["title"], 0)
        doc.add_paragraph(labels["topic"] + ": " + q)
        doc.add_picture(str(map_png), width=Inches(6.4))
        doc.add_heading(labels["analysis"], level=1)
        for block in re.split(r"\n{2,}", analysis):
            if block.strip(): doc.add_paragraph(block.strip())
        dims = (visuals.get("infographic", {}) if isinstance(visuals, dict) else {}).get("dimensions", [])
        if dims:
            doc.add_heading(labels["dims"], level=1)
            table = doc.add_table(rows=1, cols=3)
            for i, h in enumerate(["بُعد", "امتیاز کیفی", "یادداشت"]): table.rows[0].cells[i].text = h
            for x in dims[:6]:
                cells = table.add_row().cells
                cells[0].text = str(x.get("label", "")); cells[1].text = str(x.get("score", "")); cells[2].text = str(x.get("note", ""))
        if sources:
            doc.add_heading(labels["sources"], level=1)
            for i, s in enumerate(sources[:10], 1):
                doc.add_paragraph(f"{i}. {s.get('title','')} — {s.get('group','')} — {s.get('officiality','')}\n{s.get('url','')}")
        doc.save(out)
    data = base64.b64encode(out.read_bytes()).decode("ascii")
    mime = "application/pdf" if fmt == "pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return {"status":"ok", "filename":out.name, "mime":mime, "file_data":"data:"+mime+";base64,"+data, "sources":sources}


@app.get("/")
def enhanced_home():
    path = Path("dashboard.html")
    page = path.read_text(encoding="utf-8")
    panel = r'''
<section id="sourceCenter" class="analysisPanel" style="margin-top:18px">
  <div class="panelHead"><div><h2>🌐 مرکز منابع و پژوهش</h2><small>خبرگزاری‌ها، دانشگاه‌ها، اندیشکده‌ها، منابع رسمی بین‌المللی، عربی، اسرائیلی و منابع غیررسمی با ذکر منبع.</small></div></div>
  <div class="chips" style="margin-top:12px">
    <select id="sourceGroup" class="chip" style="padding:8px 12px;background:#071a2d;color:#d8e8f8;border:1px solid #254565">
      <option>همه</option>
      <option>خبرگزاری‌های داخلی</option><option>خبرگزاری‌های خارجی</option><option>دانشگاهی</option>
      <option>راهبردی و سیاسی</option><option>نظامی و امنیتی</option><option>اقتصاد و انرژی</option>
      <option>رسمی بین‌المللی</option><option>عربی</option><option>اسرائیلی</option><option>غیررسمی</option>
    </select>
    <input id="sourceQuery" class="query" style="min-height:45px;margin-top:0;flex:1" placeholder="موضوع یا کلیدواژه جستجو در منابع...">
    <button class="primary" style="width:auto" onclick="searchSources()">جستجوی منابع 🔎</button>
  </div>
  <div id="sourceStatus" class="status">منابع به‌صورت زنده از وب جستجو می‌شوند.</div>
  <div id="sourceResults" class="analysisResult" style="min-height:80px;margin-top:12px">برای شروع، موضوع را وارد کنید.</div>
</section>
<script>
async function searchSources(){
  const q=document.getElementById('sourceQuery').value.trim()||document.getElementById('query')?.value.trim();
  const group=document.getElementById('sourceGroup').value;
  const box=document.getElementById('sourceResults'); const st=document.getElementById('sourceStatus');
  if(!q){box.innerHTML='موضوع جستجو را وارد کنید.';return;}
  st.textContent='در حال جستجوی منابع...'; box.innerHTML='در حال دریافت منابع معتبر و مرتبط...';
  try{
    const r=await fetch('/api/source-search',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({query:q,group,limit:12})});
    const d=await r.json();
    if(d.status!=='ok'){box.innerHTML='<span class="error">'+(d.message||'خطا')+'</span>';st.textContent='جستجو ناموفق بود';return;}
    st.textContent='منابع پیدا شده: '+d.sources.length;
    box.innerHTML=d.sources.map((s,i)=>`<div style="padding:10px 0;border-bottom:1px solid #193a5b"><b>${i+1}. ${escapeHtml(s.title)}</b><br><small>${escapeHtml(s.group)} • ${escapeHtml(s.officiality)}</small><br><a href="${encodeURI(s.url)}" target="_blank" rel="noopener" style="color:#78baff;word-break:break-all">${escapeHtml(s.url)}</a></div>`).join('')||'منبع مستقیمی پیدا نشد.';
  }catch(e){st.textContent='خطا';box.innerHTML='<span class="error">ارتباط با موتور جستجوی منابع برقرار نشد.</span>';}
}
function escapeHtml(x){return String(x||'').replace(/[&<>'"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;',"'":'&#39;','"':'&quot;'}[c]));}
(function(){
 const oldFetch=window.fetch;
 window.fetch=async function(){
   const response=await oldFetch.apply(this,arguments);
   const url=typeof arguments[0]==='string'?arguments[0]:(arguments[0]?.url||'');
   if(url.endsWith('/api/analyze')){
     try{const clone=response.clone();const data=await clone.json();if(data.sources?.length){const box=document.getElementById('sourceResults');if(box){box.innerHTML='<h3>منابع مورد استفاده در تحلیل</h3>'+data.sources.map((s,i)=>`<div style="padding:8px 0;border-bottom:1px solid #193a5b"><b>${i+1}. ${escapeHtml(s.title)}</b><br><small>${escapeHtml(s.group)} • ${escapeHtml(s.officiality)}</small><br><a href="${encodeURI(s.url)}" target="_blank" rel="noopener" style="color:#78baff">${escapeHtml(s.url)}</a></div>`).join('');}}}catch(e){}
   }
   return response;
 };
})();
</script>
'''
    return HTMLResponse(page.replace("</body>", panel + "</body>"))


@app.get("/api/sources")
def source_catalog():
    return {"status":"ok", "groups": SOURCE_GROUPS, "domains": ALL_DOMAINS}


@app.post("/api/source-search")
def source_search(request: SourceSearchRequest):
    return search_web(request.query, request.group, request.limit)


@app.post("/api/analyze")
def enhanced_analyze(request: AnalysisRequest):
    result = core_analyze(request)
    if result.get("status") != "ok":
        return result
    search = search_web(request.query, "همه", 10)
    result["sources"] = search.get("sources", [])
    result["source_status"] = search.get("status")
    return result


@app.post("/api/export-report")
def enhanced_export(request: ExportRequest):
    try:
        sources = sources_for_report(request.query, request.analysis)
        return build_source_report(request, sources)
    except Exception as exc:
        logger.exception("SOURCE_REPORT_EXPORT_ERROR")
        return {"status":"error", "message":"ساخت گزارش با منابع ناموفق بود: "+(str(exc).strip()[:500] or "خطای نامشخص")}


@app.post("/api/publish-text")
def enhanced_publish(request: PublishRequest):
    sources = sources_for_report(request.query, request.analysis)
    lines = ["📌 گزارش تحلیلی — AI Strategic Studio", "", "موضوع: " + request.query.strip(), "", "🧠 جمع‌بندی", clean_text(request.analysis)[:6500]]
    if sources:
        lines += ["", "🔎 منابع مرتبط"]
        lines += [f"• {s.get('title','')} — {s.get('group','')} — {s.get('url','')}" for s in sources[:8]]
    lines += ["", "این متن برای انتشار در شبکه‌های اجتماعی آماده شده و ادعاهای حساس باید پیش از انتشار راستی‌آزمایی شوند."]
    return {"status":"ok", "text":"\n".join(lines), "sources":sources}


# Keep every existing dashboard/API route while adding the research gateway above it.
app.mount("/", core_app)
